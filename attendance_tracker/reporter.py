from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from dateutil import parser


def clean_display_time(val):
    if not val or val == "-":
        return "-"
    try:
        dt = parser.isoparse(str(val))
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(val)


def clean_display_email(val):
    if not val or val == "-":
        return "-"
    val = str(val).strip()
    if val.startswith("users/"):
        uid = val.replace("users/", "")
        short_id = f"{uid[:4]}...{uid[-4:]}" if len(uid) > 8 else uid
        return f"Google Account ({short_id})"
    return val


class AttendanceReporter:
    """
    Generates styled Excel and CSV reports with visual status badges,
    attendance summaries, and clean layouts.
    """

    @staticmethod
    def export_excel(
        attendance_data: List[Dict[str, Any]],
        output_path: Path,
        meeting_title: str = "Google Meet Attendance",
        meeting_date: str = ""
    ) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Summary"
        ws.views.sheetView[0].showGridLines = True

        if not meeting_date:
            meeting_date = datetime.now().strftime("%Y-%m-%d %H:%M")

        # Header Styles
        title_font = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        table_header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        table_header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")

        # Metric Card Styles
        card_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        card_font_label = Font(name="Segoe UI", size=9, color="64748B")
        card_font_val = Font(name="Segoe UI", size=14, bold=True, color="0F172A")

        # Status Styles
        present_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        present_font = Font(name="Segoe UI", size=10, bold=True, color="15803D")
        absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        absent_font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
        data_font = Font(name="Segoe UI", size=10, color="1E293B")

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # 1. Title Section
        ws["A1"] = meeting_title
        ws["A1"].font = title_font
        ws["A2"] = f"Generated on: {meeting_date}"
        ws["A2"].font = sub_font

        # 2. Summary Metrics
        total_attendees = len(attendance_data)
        present_count = sum(1 for r in attendance_data if r.get("status") == "Present")
        absent_count = total_attendees - present_count
        attendance_rate = (present_count / total_attendees * 100) if total_attendees > 0 else 0

        ws["A4"] = "Total Attendees"
        ws["A5"] = total_attendees
        ws["B4"] = "Present"
        ws["B5"] = present_count
        ws["C4"] = "Absent / Short"
        ws["C5"] = absent_count
        ws["D4"] = "Attendance %"
        ws["D5"] = f"{attendance_rate:.1f}%"

        for col in ["A", "B", "C", "D"]:
            c4 = ws[f"{col}4"]
            c5 = ws[f"{col}5"]
            c4.font = card_font_label
            c4.fill = card_fill
            c4.alignment = Alignment(horizontal="center", vertical="center")
            c5.font = card_font_val
            c5.fill = card_fill
            c5.alignment = Alignment(horizontal="center", vertical="center")
            c4.border = thin_border
            c5.border = thin_border

        # 3. Table Headers
        has_roll_no = any(bool(r.get("roll_no")) for r in attendance_data)
        has_real_email = any(r.get("email") and "@" in str(r.get("email")) for r in attendance_data)

        headers = ["Roll No." if has_roll_no else "S.No.", "Participant Name"]
        if has_real_email:
            headers.append("Email")
        headers.extend(["Duration (min)", "Sessions", "Status", "Join Time", "Leave Time"])
        start_row = 7

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = table_header_font
            cell.fill = table_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # 4. Table Data Rows
        current_row = start_row + 1
        for idx, record in enumerate(attendance_data, start=1):
            c_idx = 1
            first_val = record.get("roll_no", "") if has_roll_no else idx
            ws.cell(row=current_row, column=c_idx, value=first_val).alignment = Alignment(horizontal="center")
            c_idx += 1

            ws.cell(row=current_row, column=c_idx, value=record.get("name", "")).alignment = Alignment(horizontal="left")
            c_idx += 1

            if has_real_email:
                ws.cell(row=current_row, column=c_idx, value=record.get("email", "-")).alignment = Alignment(horizontal="left")
                c_idx += 1

            ws.cell(row=current_row, column=c_idx, value=record.get("duration_minutes", 0)).alignment = Alignment(horizontal="right")
            c_idx += 1

            ws.cell(row=current_row, column=c_idx, value=record.get("session_count", 1)).alignment = Alignment(horizontal="center")
            c_idx += 1

            status_cell = ws.cell(row=current_row, column=c_idx, value=record.get("status", ""))
            status_cell.alignment = Alignment(horizontal="center")
            if record.get("status") == "Present":
                status_cell.fill = present_fill
                status_cell.font = present_font
            else:
                status_cell.fill = absent_fill
                status_cell.font = absent_font
            status_col_num = c_idx
            c_idx += 1

            ws.cell(row=current_row, column=c_idx, value=clean_display_time(record.get("join_time", "-"))).alignment = Alignment(horizontal="center")
            c_idx += 1

            ws.cell(row=current_row, column=c_idx, value=clean_display_time(record.get("leave_time", "-"))).alignment = Alignment(horizontal="center")

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=current_row, column=col_idx)
                if col_idx != status_col_num:
                    c.font = data_font
                c.border = thin_border

            current_row += 1

        # Adjust column widths automatically with generous width for Name
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            # Give participant name (Column B) generous width
            if col_letter == "B":
                ws.column_dimensions[col_letter].width = max(max_len + 4, 30)
            else:
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        wb.save(output_path)
        return output_path

    @staticmethod
    def export_csv(attendance_data: List[Dict[str, Any]], output_path: Path) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        df = pd.DataFrame(attendance_data)
        df.to_csv(output_path, index=False)
        return output_path
